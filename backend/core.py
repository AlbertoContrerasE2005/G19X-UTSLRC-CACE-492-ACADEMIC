"""Motor DataOps: parseo, validación, IA y worker.

- Acepta CSV / XLSX / SQL con 1-30 columnas cualquiera.
- Ventas (5 columnas exactas): reglas estrictas + IA con referencia.
- Dinámico: validación genérica + IA sobre columnas numéricas.
"""

import csv
import hashlib
import io
import json
import math
import re
import secrets
import threading
import time
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import numpy as np
from sklearn.ensemble import IsolationForest

from backend import db as mydb

ROOT = Path(__file__).resolve().parents[1]

# Columnas del caso ventas (único caso con reglas estrictas).
COLUMNS = ["id", "fecha", "producto", "cantidad", "precio_unitario"]

MAX_BYTES = 10 * 1024 * 1024  # 10 MiB por archivo.
MAX_ROWS = 20000  # Filas máximas por archivo.
MODEL_SALES = "ventas-v1"
MODEL_GENERIC = "generico-v1"


# ---------------------------------------------------------------- utilidades


def now():
    """Fecha actual ISO UTC."""
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


def clean_header(name):
    """Normaliza un nombre de columna: minúsculas, sin espacios raros."""
    name = str(name or "").strip().lower().replace(" ", "_")[:64]
    name = re.sub(r"[^a-z0-9_áéíóúñ.\-]", "", name)
    return name or "col"


def password_hash(password):
    """Hash scrypt con sal de 16 bytes. Formato sal:hash en hex."""
    salt = secrets.token_bytes(16)
    value = hashlib.scrypt(password.encode(), salt=salt, n=16384, r=8, p=1)
    return salt.hex() + ":" + value.hex()


def password_ok(password, stored):
    """Compara contraseña contra hash (tiempo constante)."""
    salt, expected = stored.split(":")
    value = hashlib.scrypt(
        password.encode(), salt=bytes.fromhex(salt), n=16384, r=8, p=1
    )
    return secrets.compare_digest(value.hex(), expected)


# ---------------------------------------------------------------- parseo


def parse_csv_dynamic(content):
    """Lee CSV (coma o punto y coma) con cualquier columna.

    Devuelve (columns, rows). Falla si: vacío, >10 MiB, binario,
    cabecera inválida, filas disparejas, >20000 filas.
    """
    if not isinstance(content, str) or not content.strip():
        raise ValueError("El archivo está vacío.")
    if len(content.encode("utf-8")) > MAX_BYTES:
        raise ValueError("El tamaño máximo es 10 MiB.")
    if "\x00" in content:
        raise ValueError("El archivo tiene caracteres binarios. Usa UTF-8.")

    content = content.lstrip("\ufeff")
    header = content.splitlines()[0]
    delimiter = ";" if header.count(";") > header.count(",") else ","

    try:
        reader = csv.DictReader(
            io.StringIO(content, newline=""), delimiter=delimiter, strict=True
        )
        raw = reader.fieldnames or []
        columns = [clean_header(c) for c in raw]
        if (
            not columns
            or len(columns) > 30
            or any(not c for c in columns)
            or len(set(columns)) != len(columns)
        ):
            raise ValueError("Cabecera inválida: 1 a 30 columnas únicas.")

        rows = []
        for item in reader:
            if len(rows) >= MAX_ROWS:
                raise ValueError("El límite es de 20000 registros.")
            if None in item or any(v is None for v in item.values()):
                raise ValueError("Hay filas con columnas de más o de menos.")
            row = {
                col: str(item[orig]).strip() if item[orig] is not None else ""
                for col, orig in zip(columns, raw)
            }
            if any(len(v) > 2000 for v in row.values()):
                raise ValueError("Cada celda admite hasta 2000 caracteres.")
            rows.append(row)
    except csv.Error as exc:
        raise ValueError("CSV con comillas o separadores incorrectos.") from exc

    if not rows:
        raise ValueError("El archivo debe tener al menos un registro.")
    return columns, rows


def parse_csv_strict(content):
    """Valida que el CSV sea exactamente la plantilla de ventas."""
    columns, rows = parse_csv_dynamic(content)
    if columns != COLUMNS:
        raise ValueError(
            "Para ventas, las columnas deben ser: " + ", ".join(COLUMNS) + "."
        )
    return rows


def parse_excel_dynamic(data):
    """Lee XLSX (primera hoja) con cualquier columna."""
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise ValueError("El Excel está vacío.")
    if len(data) > MAX_BYTES:
        raise ValueError("El Excel supera los 10 MiB.")

    from openpyxl import load_workbook

    try:
        wb = load_workbook(
            filename=io.BytesIO(bytes(data)), read_only=True, data_only=True
        )
        values = list(wb.active.iter_rows(values_only=True))
    except Exception as exc:
        raise ValueError("No se pudo leer el Excel.") from exc

    if not values or len(values) < 2:
        raise ValueError("El Excel necesita cabecera + 1 registro mínimo.")

    raw = [str(c).strip() if c is not None else "" for c in values[0]]
    columns = [clean_header(c) for c in raw]
    if (
        not columns
        or len(columns) > 30
        or any(not c for c in columns)
        or len(set(columns)) != len(columns)
    ):
        raise ValueError("Cabecera Excel inválida: 1 a 30 columnas únicas.")

    rows = []
    for line in values[1:]:
        if all(c is None or str(c).strip() == "" for c in line):
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError("El límite es de 20000 registros.")
        cells = [("" if c is None else str(c).strip()) for c in list(line)]
        while len(cells) < len(columns):
            cells.append("")
        rows.append(dict(zip(columns, cells[: len(columns)])))

    if not rows:
        raise ValueError("El archivo debe tener al menos un registro.")
    return columns, rows


def parse_sql_dynamic(content):
    """Lee .sql con INSERT INTO tabla (col1, col2) VALUES (...).

    Acepta cualquier columna. Ignora comentarios -- y /* */.
    """
    if not isinstance(content, str) or not content.strip():
        raise ValueError("El SQL está vacío.")
    if len(content.encode("utf-8")) > MAX_BYTES:
        raise ValueError("El tamaño máximo es 10 MiB.")

    text = re.sub(r"/\*.*?\*/", " ", content, flags=re.DOTALL)
    kept = []
    for line in text.splitlines():
        s = line.strip()
        if not s or s.startswith("--"):
            continue
        kept.append(re.split(r"\s--\s", line)[0])
    text = " ".join(kept)

    inserts = list(
        re.finditer(
            r"INSERT\s+INTO\s+\S+\s*\(([^)]+)\)\s*VALUES\s*(.+?);",
            text,
            flags=re.IGNORECASE | re.DOTALL,
        )
    )
    if not inserts:
        raise ValueError(
            "Sin INSERTs. Usa: INSERT INTO tabla (col1, col2) VALUES (...);"
        )

    def split_cols(s):
        return [
            clean_header(c.strip().strip('"').strip("'").strip("`").strip("[]"))
            for c in s.split(",")
        ]

    columns = split_cols(inserts[0].group(1))
    if not columns or len(columns) > 30 or len(set(columns)) != len(columns):
        raise ValueError("Columnas SQL inválidas: 1 a 30 únicas.")

    token = re.compile(r"""'((?:[^']|'')*)'|"((?:[^"]|"")*)"|([^,()]+)""")
    rows = []
    for ins in inserts:
        cols_here = split_cols(ins.group(1))
        for m in re.finditer(r"\(([^;()]+)\)", ins.group(2)):
            cells = []
            for tok in token.finditer(m.group(1)):
                if tok.group(1) is not None:
                    cells.append(tok.group(1).replace("''", "'").strip())
                elif tok.group(2) is not None:
                    cells.append(tok.group(2).replace('""', '"').strip())
                else:
                    cells.append((tok.group(3) or "").strip())
            if len(cells) != len(cols_here):
                continue
            row = dict(zip(cols_here, cells))
            rows.append({c: row.get(c, "") for c in columns})
            if len(rows) >= MAX_ROWS:
                raise ValueError("El límite es de 20000 registros.")

    if not rows:
        raise ValueError("El SQL no aportó filas válidas.")
    return columns, rows


def rows_to_csv(columns, rows):
    """Convierte filas a texto CSV (formato interno de storage)."""
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=columns)
    writer.writeheader()
    for r in rows:
        writer.writerow({k: r.get(k, "") for k in columns})
    return out.getvalue()


# ---------------------------------------------------------------- validación


def validate_sales(rows):
    """Reglas estrictas de ventas. Calcula total y features para IA."""
    from datetime import date

    results, seen = [], set()
    for number, raw in enumerate(rows, 2):
        data = {k: str(raw[k]).strip() for k in COLUMNS}
        errors = []

        ident = data["id"]
        if not ident or len(ident) > 64:
            errors.append("ID obligatorio hasta 64 caracteres")
        if ident in seen:
            errors.append("ID repetido en el archivo")
        seen.add(ident)

        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", data["fecha"]):
            errors.append("Fecha formato AAAA-MM-DD")
        else:
            try:
                date.fromisoformat(data["fecha"])
            except ValueError:
                errors.append("Fecha inexistente")

        if not data["producto"] or len(data["producto"]) > 120:
            errors.append("Producto obligatorio hasta 120 caracteres")

        try:
            qty = Decimal(data["cantidad"])
            if (
                not qty.is_finite()
                or qty != qty.to_integral_value()
                or not 0 < qty <= 1_000_000_000
            ):
                raise ValueError()
        except (ValueError, InvalidOperation):
            errors.append("Cantidad entera 1 a 1000000000")

        try:
            price = Decimal(data["precio_unitario"])
            if (
                not price.is_finite()
                or not 0 <= price <= 1_000_000_000_000
                or price.as_tuple().exponent < -2
            ):
                raise ValueError()
        except (ValueError, InvalidOperation):
            errors.append("Precio 0 a 1e12 con 2 decimales máximo")

        item = {
            "line": number,
            "data": data,
            "status": "invalid" if errors else "valid",
            "reason": "; ".join(errors),
            "score": None,
        }
        if not errors:
            item["features"] = [float(qty), float(price), float(qty * price)]
            data["cantidad"] = str(int(qty))
            data["precio_unitario"] = str(price.quantize(Decimal(".01")))
            data["total"] = str((qty * price).quantize(Decimal(".01")))
        results.append(item)
    return results


def validate_generic(columns, rows):
    """Validación sin plantilla: primera columna = clave única obligatoria."""
    results, seen = [], set()
    key = columns[0]
    for number, raw in enumerate(rows, 2):
        data = {k: str(raw.get(k, "")).strip() for k in columns}
        errors = []
        ident = data.get(key, "")
        if not ident or len(ident) > 256:
            errors.append(f"{key} obligatorio hasta 256 caracteres")
        if ident in seen:
            errors.append(f"{key} repetido en el archivo")
        seen.add(ident)
        if all(v == "" for v in data.values()):
            errors.append("Fila vacía")
        results.append(
            {
                "line": number,
                "data": data,
                "status": "invalid" if errors else "valid",
                "reason": "; ".join(errors),
                "score": None,
            }
        )
    return results


# ---------------------------------------------------------------- storage


class Store:
    """Acceso MySQL. Cada método abre y cierra su conexión."""

    def one(self, sql, params=()):
        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                cur.execute(sql, params)
                return cur.fetchone()
        finally:
            conn.close()

    def all(self, sql, params=()):
        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                cur.execute(sql, params)
                return list(cur.fetchall())
        finally:
            conn.close()

    def run(self, sql, params=()):
        """INSERT/UPDATE/DELETE. Devuelve (id_insertado, filas_afectadas)."""
        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                n = cur.execute(sql, params)
                return cur.lastrowid, n
        finally:
            conn.close()

    # -- helpers de dominio -------------------------------------------

    def audit(self, user_id, action, entity=None):
        self.run(
            "INSERT INTO audit (created_at, user_id, action, entity_id)"
            " VALUES (%s, %s, %s, %s)",
            (now(), user_id, action, entity),
        )

    def activity(
        self,
        user_id,
        username,
        action,
        entity_type="",
        entity_id="",
        project_id=None,
        detail="",
        result="",
    ):
        """Historial visible de actividad (quien, qué, dónde, cuándo)."""
        self.run(
            "INSERT INTO activity_logs (created_at, user_id, username, action,"
            " entity_type, entity_id, project_id, detail, result)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
            (
                now(),
                user_id,
                username or "",
                action,
                entity_type,
                str(entity_id or ""),
                project_id,
                detail or "",
                result or "",
            ),
        )

    def log(self, run_id, stage, message):
        self.run(
            "INSERT INTO run_logs (run_id, created_at, stage, message)"
            " VALUES (%s, %s, %s, %s)",
            (run_id, now(), stage, message),
        )
        self.run("UPDATE runs SET stage = %s WHERE id = %s", (stage, run_id))

    def alert(
        self,
        level,
        title,
        message,
        entity_type="",
        entity_id="",
        project_id=None,
        user_id=None,
    ):
        """Crea una alerta visible en Monitoreo/Alertas."""
        new_id, _ = self.run(
            "INSERT INTO alerts (project_id, user_id, level, title, message,"
            " entity_type, entity_id, `read`, created_at)"
            " VALUES (%s, %s, %s, %s, %s, %s, %s, 0, %s)",
            (
                project_id,
                user_id,
                level,
                title[:200],
                message[:1000],
                entity_type,
                entity_id,
                now(),
            ),
        )
        return new_id

    def queue(
        self, dataset_id, user_id, kind="manual", parent=None, pipeline="p_ventas"
    ):
        """Encola una ejecución. Solo 1 activa global (compatibilidad)."""
        active = self.one(
            "SELECT id FROM runs WHERE status IN ('queued', 'running') LIMIT 1"
        )
        if active:
            raise ValueError("Ya hay una ejecución en curso. Espera.")
        if not self.one("SELECT id FROM pipelines WHERE id = %s", (pipeline,)):
            raise ValueError("Pipeline no encontrado.")
        run_id = uuid.uuid4().hex
        self.run(
            "INSERT INTO runs (id, dataset_id, pipeline_id, trigger_kind,"
            " status, created_at, created_by, parent_id)"
            " VALUES (%s, %s, %s, %s, 'queued', %s, %s, %s)",
            (run_id, dataset_id, pipeline, kind, now(), user_id, parent),
        )
        self.audit(user_id, "run_queued", run_id)
        return run_id

    def review(self, run_id, lines, user_id, username, motivo, action):
        """Aprueba (carga) o rechaza anomalías pendientes."""
        if action not in ("approved", "rejected"):
            raise ValueError("Acción inválida.")
        if not isinstance(lines, list) or not 1 <= len(lines) <= 500:
            raise ValueError("Selecciona de 1 a 500 filas.")
        try:
            ids = sorted({int(x) for x in lines})
        except (ValueError, TypeError):
            raise ValueError("Filas inválidas.")
        motivo = str(motivo or "").strip()[:500]
        if len(motivo) < 3:
            raise ValueError("Motivo de al menos 3 caracteres.")

        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT r.*, d.columns_json FROM runs r"
                    " JOIN datasets d ON d.id = r.dataset_id"
                    " WHERE r.id = %s AND r.status = 'completed'",
                    (run_id,),
                )
                run = cur.fetchone()
                if not run:
                    raise ValueError("Solo runs completados se revisan.")
                columns = json.loads(run["columns_json"] or "[]") or COLUMNS
                is_sales = columns == COLUMNS

                marks = ",".join(["%s"] * len(ids))
                cur.execute(
                    f"SELECT line, data_json FROM run_rows"
                    f" WHERE run_id = %s AND line IN ({marks})"
                    f" AND status = 'anomaly'",
                    (run_id, *ids),
                )
                pending = cur.fetchall()
                if not pending:
                    raise ValueError("Sin anomalías pendientes en esas filas.")

                loaded = existing = rejected = 0
                for row in pending:
                    d = json.loads(row["data_json"])
                    if action == "approved":
                        if is_sales:
                            done = cur.execute(
                                "INSERT IGNORE INTO sales"
                                " (id, date, product, quantity,"
                                " unit_price, total, source_run)"
                                " VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                (
                                    d["id"],
                                    d["fecha"],
                                    d["producto"],
                                    int(d["cantidad"]),
                                    d["precio_unitario"],
                                    d.get("total", "0.00"),
                                    run_id,
                                ),
                            )
                        else:
                            key = str(d.get(columns[0], ""))[:256]
                            done = cur.execute(
                                "INSERT IGNORE INTO records"
                                " (pipeline_id, run_id, `key`, data_json, created_at)"
                                " VALUES (%s, %s, %s, %s, %s)",
                                (
                                    run["pipeline_id"],
                                    run_id,
                                    key,
                                    json.dumps(d, ensure_ascii=False),
                                    now(),
                                ),
                            )
                        if done:
                            status, reason, loaded = (
                                "loaded",
                                f"Aprobado por {username}: {motivo}",
                                loaded + 1,
                            )
                        else:
                            status, reason, existing = (
                                "existing",
                                f"Aprobado pero ya existía ({username}): {motivo}",
                                existing + 1,
                            )
                    else:
                        status, reason, rejected = (
                            "rejected",
                            f"Rechazado por {username}: {motivo}",
                            rejected + 1,
                        )
                    cur.execute(
                        "UPDATE run_rows SET status = %s, reason = %s"
                        " WHERE run_id = %s AND line = %s",
                        (status, reason, run_id, row["line"]),
                    )
                    cur.execute(
                        "INSERT INTO approvals"
                        " (run_id, line, action, motivo, user_id, created_at)"
                        " VALUES (%s, %s, %s, %s, %s, %s)",
                        (run_id, row["line"], action, motivo, user_id, now()),
                    )

                n = len(pending)
                if action == "approved":
                    cur.execute(
                        "UPDATE runs SET anomalies = anomalies - %s,"
                        " loaded = loaded + %s, existing = existing + %s,"
                        " approved = approved + %s WHERE id = %s",
                        (n, loaded, existing, n, run_id),
                    )
                else:
                    cur.execute(
                        "UPDATE runs SET anomalies = anomalies - %s,"
                        " rejected = rejected + %s WHERE id = %s",
                        (n, n, run_id),
                    )
                cur.execute(
                    "INSERT INTO run_logs (run_id, created_at, stage, message)"
                    " VALUES (%s, %s, 'approval', %s)",
                    (run_id, now(), f"{n} {action} por {username}: {motivo}"),
                )
                cur.execute(
                    "INSERT INTO audit (created_at, user_id, action, entity_id)"
                    " VALUES (%s, %s, %s, %s)",
                    (now(), user_id, f"anomaly_{action}", run_id),
                )
        finally:
            conn.close()
        return {
            "reviewed": n,
            "loaded": loaded,
            "existing": existing,
            "rejected": rejected,
        }


# ---------------------------------------------------------------- motor


class Engine:
    """Procesa un run: valida → IA → carga. Se entrena al iniciar."""

    def __init__(self, store):
        self.store = store
        ref = (ROOT / "examples" / "referencia_ventas.csv").read_text(encoding="utf-8")
        _, rows = parse_csv_dynamic(ref)
        feats = [r["features"] for r in validate_sales(rows) if r["status"] == "valid"]
        self.model = IsolationForest(
            n_estimators=120, contamination="auto", random_state=42, n_jobs=1
        )
        self.model.fit(np.log1p(np.array(feats)))

    def process(self, run_id):
        start = time.perf_counter()
        rec = self.store.one(
            "SELECT r.*, d.content, d.columns_json FROM runs r"
            " JOIN datasets d ON d.id = r.dataset_id WHERE r.id = %s",
            (run_id,),
        )
        columns = json.loads(rec["columns_json"] or "[]") or COLUMNS
        is_sales = columns == COLUMNS
        model = MODEL_SALES

        # 1. Validación.
        if is_sales:
            self.store.log(run_id, "validation", "Reglas de ventas.")
            _, sales_rows = parse_csv_dynamic(rec["content"])
            rows = validate_sales(sales_rows)
        else:
            self.store.log(
                run_id,
                "validation",
                f"Esquema dinámico ({len(columns)} columnas).",
            )
            _, dyn = parse_csv_dynamic(rec["content"])
            rows = validate_generic(
                columns, [{c: r.get(c, "") for c in columns} for r in dyn]
            )
        valid = [r for r in rows if r["status"] == "valid"]

        # 2. Anomalías con IA.
        if is_sales:
            self.store.log(
                run_id,
                "anomaly",
                f"{len(valid)} válidos. Modelo {MODEL_SALES}.",
            )
            if valid:
                scores = self.model.decision_function(
                    np.log1p(np.array([r["features"] for r in valid]))
                )
                for item, s in zip(valid, scores):
                    item["score"] = float(s)
                    if s < 0:
                        item["status"] = "anomaly"
                        item["reason"] = "Atípico según IA. Revisar."
        else:
            nums = []
            for c in columns[1:]:
                try:
                    vals = [
                        float(str(r["data"].get(c, "")).replace(",", "")) for r in valid
                    ]
                    if all(math.isfinite(v) for v in vals) and len(set(vals)) > 1:
                        nums.append(c)
                except (ValueError, TypeError):
                    continue
            if valid and nums and len(valid) >= 3:
                try:
                    matrix = np.array(
                        [
                            [
                                float(str(r["data"].get(c, "0")).replace(",", ""))
                                for c in nums
                            ]
                            for r in valid
                        ]
                    )
                    m = IsolationForest(
                        n_estimators=100,
                        contamination="auto",
                        random_state=42,
                        n_jobs=1,
                    )
                    for item, s in zip(valid, m.fit(matrix).decision_function(matrix)):
                        item["score"] = float(s)
                        if s < 0:
                            item["status"] = "anomaly"
                            item["reason"] = f"Atípico en {', '.join(nums)}."
                    model = MODEL_GENERIC
                    self.store.log(
                        run_id,
                        "anomaly",
                        f"IA genérica en {', '.join(nums)}:"
                        f" {sum(r['status'] == 'anomaly' for r in valid)} alertas.",
                    )
                except Exception:
                    self.store.log(run_id, "anomaly", "IA no aplicable.")
            else:
                self.store.log(run_id, "anomaly", "Sin IA aplicable.")

        # 3. Carga (válidos sin alerta; repetidos no duplican).
        self.store.log(run_id, "loading", "Cargando a destino.")
        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM run_rows WHERE run_id = %s", (run_id,))
                pipe = rec["pipeline_id"] or "p_ventas"
                loaded = existing = 0
                for item in rows:
                    if item["status"] == "valid":
                        d = item["data"]
                        if is_sales:
                            done = cur.execute(
                                "INSERT IGNORE INTO sales (id, date, product,"
                                " quantity, unit_price, total, source_run)"
                                " VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                (
                                    d["id"],
                                    d["fecha"],
                                    d["producto"],
                                    int(d["cantidad"]),
                                    d["precio_unitario"],
                                    d["total"],
                                    run_id,
                                ),
                            )
                        else:
                            done = cur.execute(
                                "INSERT IGNORE INTO records (pipeline_id, run_id,"
                                " `key`, data_json, created_at)"
                                " VALUES (%s, %s, %s, %s, %s)",
                                (
                                    pipe,
                                    run_id,
                                    str(d.get(columns[0], ""))[:256],
                                    json.dumps(d, ensure_ascii=False),
                                    now(),
                                ),
                            )
                        if done:
                            item["status"], loaded = "loaded", loaded + 1
                        else:
                            item["status"] = "existing"
                            item["reason"] = "Clave ya existente; se conserva."
                            existing += 1
                    cur.execute(
                        "INSERT INTO run_rows"
                        " (run_id, line, data_json, status, reason, score)"
                        " VALUES (%s, %s, %s, %s, %s, %s)",
                        (
                            run_id,
                            item["line"],
                            json.dumps(item["data"], ensure_ascii=False),
                            item["status"],
                            item["reason"],
                            item.get("score"),
                        ),
                    )
                invalid = sum(r["status"] == "invalid" for r in rows)
                anomalies = sum(r["status"] == "anomaly" for r in rows)
                ms = round((time.perf_counter() - start) * 1000)
                cur.execute(
                    "UPDATE runs SET status = 'completed', stage = 'done',"
                    " finished_at = %s, total = %s, valid = %s, invalid = %s,"
                    " anomalies = %s, loaded = %s, existing = %s,"
                    " model_version = %s, duration_ms = %s, error = NULL"
                    " WHERE id = %s",
                    (
                        now(),
                        len(rows),
                        len(valid),
                        invalid,
                        anomalies,
                        loaded,
                        existing,
                        model,
                        ms,
                        run_id,
                    ),
                )
                cur.execute(
                    "INSERT INTO run_logs (run_id, created_at, stage, message)"
                    " VALUES (%s, %s, 'done', %s)",
                    (
                        run_id,
                        now(),
                        f"{loaded} cargados, {invalid} inválidos,"
                        f" {anomalies} alertas, {existing} existentes.",
                    ),
                )
                cur.execute(
                    "UPDATE pipelines SET last_run = %s, run_count = run_count + 1,"
                    " status = 'activo' WHERE id = %s",
                    (now(), pipe),
                )
                # 4. Post-proceso: calidad, reglas, anomalías y alertas.
                try:
                    from backend.profiling import (
                        evaluate_rule,
                        quality_score,
                    )

                    ds = self.store.one(
                        "SELECT project_id FROM datasets WHERE id = %s",
                        (rec["dataset_id"],),
                    )
                    proj = (ds and ds.get("project_id")) or "p_default"
                    raw = [dict(it["data"]) for it in rows]
                    qual = quality_score(columns, raw)
                    self.store.run(
                        "UPDATE datasets SET quality_score = %s WHERE id = %s",
                        (qual.get("general", 0), rec["dataset_id"]),
                    )
                    self.store.run(
                        "DELETE FROM quality_results WHERE run_id = %s", (run_id,)
                    )
                    for issue in (qual.get("issues") or [])[:30]:
                        sev = str(issue.get("severity", "media")).lower()
                        if sev not in ("baja", "media", "alta"):
                            sev = "media"
                        self.store.run(
                            "INSERT INTO quality_results (run_id, dataset_id,"
                            " project_id, check_type, column_name, message,"
                            " severity, affected, created_at)"
                            " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
                            (
                                run_id,
                                rec["dataset_id"],
                                proj,
                                "automatica",
                                str(issue.get("column", ""))[:120],
                                str(issue.get("problem", ""))[:500],
                                sev,
                                int(issue.get("affected", 0) or 0),
                                now(),
                            ),
                        )
                    for rl in self.store.all(
                        "SELECT * FROM quality_rules WHERE active = 1"
                        " AND (dataset_id = %s OR pipeline_id = %s)",
                        (rec["dataset_id"], pipe),
                    ):
                        try:
                            res = evaluate_rule(
                                rl["column_name"],
                                rl["condition"],
                                rl["value"] or "",
                                columns,
                                raw,
                            )
                        except Exception:
                            continue
                        if res.get("failed"):
                            sev = (
                                rl["severity"]
                                if rl["severity"]
                                in ("baja", "media", "alta", "critica")
                                else "media"
                            )
                            self.store.run(
                                "INSERT INTO quality_results (run_id, dataset_id,"
                                " project_id, check_type, column_name, message,"
                                " severity, affected, created_at)"
                                " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
                                (
                                    run_id,
                                    rec["dataset_id"],
                                    proj,
                                    "regla:" + rl["id"],
                                    rl["column_name"],
                                    f"Regla incumplida: {rl['column_name']}"
                                    f" {rl['condition']} {rl['value']}"
                                    f" ({res['failed']} filas).",
                                    sev,
                                    int(res["failed"]),
                                    now(),
                                ),
                            )
                    self.store.run("DELETE FROM anomalies WHERE run_id = %s", (run_id,))
                    for it in rows:
                        if it["status"] != "anomaly":
                            continue
                        score = it.get("score")
                        sev = "alta" if score is not None and score < -0.15 else "media"
                        self.store.run(
                            "INSERT INTO anomalies (run_id, dataset_id, project_id,"
                            " column_name, row_line, value, reason, severity,"
                            " score, status, created_at)"
                            " VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s,"
                            " 'pendiente', %s)",
                            (
                                run_id,
                                rec["dataset_id"],
                                proj,
                                columns[0],
                                int(it.get("line", 0)),
                                json.dumps(it["data"], ensure_ascii=False)[:500],
                                str(it.get("reason", ""))[:500],
                                sev,
                                score,
                                now(),
                            ),
                        )
                    finished = now()
                    if invalid + anomalies > 0 and qual.get("general", 100) < 90:
                        self.store.alert(
                            (
                                "critico"
                                if qual.get("general", 100) < 70
                                else "advertencia"
                            ),
                            "Calidad inferior al mínimo",
                            f"Calidad {qual.get('general')}%. {invalid} inválidos"
                            f" y {anomalies} anomalías.",
                            "run",
                            run_id,
                            proj,
                        )
                    if anomalies > 0:
                        self.store.alert(
                            "advertencia" if anomalies < 5 else "error",
                            "Anomalías detectadas",
                            f"{anomalies} valores atípicos requieren revisión.",
                            "run",
                            run_id,
                            proj,
                        )
                    if invalid > 0 and invalid / max(1, len(rows)) > 0.2:
                        self.store.alert(
                            "error",
                            "Muchos datos inválidos",
                            f"{invalid} de {len(rows)} no cumplen reglas.",
                            "run",
                            run_id,
                            proj,
                        )
                except Exception:
                    import logging

                    logging.exception("Post-proceso falló para run %s", run_id)
        finally:
            conn.close()


# ---------------------------------------------------------------- worker


class Worker:
    """Hilo único: programa pipelines y ejecuta runs encolados."""

    def __init__(self, store):
        self.store = store
        self.engine = Engine(store)
        self.stop_event = threading.Event()
        self.thread = None

    def recover(self):
        """Al arrancar: lo 'running' vuelve a 'queued'."""
        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT id FROM runs WHERE status = 'running'")
                ids = [r["id"] for r in cur.fetchall()]
                cur.execute(
                    "UPDATE runs SET status = 'queued', stage = 'queued'"
                    " WHERE status = 'running'"
                )
                cur.execute(
                    "DELETE FROM sessions WHERE expires_at < %s", (time.time(),)
                )
        finally:
            conn.close()
        for rid in ids:
            self.store.log(rid, "queued", "Recuperada tras reinicio.")

    def schedule(self):
        """Dispara pipelines vencidos (1 por ciclo, si no hay activa)."""
        if self.store.one(
            "SELECT id FROM runs WHERE status IN ('queued', 'running') LIMIT 1"
        ):
            return
        due = self.store.all(
            "SELECT * FROM pipelines WHERE enabled = 1"
            " AND dataset_id IS NOT NULL AND next_run IS NOT NULL"
            " AND next_run <= %s ORDER BY next_run LIMIT 1",
            (time.time(),),
        )
        for p in due:
            new_id = uuid.uuid4().hex
            conn = mydb.connect()
            try:
                with conn.cursor() as cur:
                    cur.execute(
                        "INSERT INTO runs (id, dataset_id, pipeline_id,"
                        " trigger_kind, status, created_at)"
                        " VALUES (%s, %s, %s, 'scheduled', 'queued', %s)",
                        (new_id, p["dataset_id"], p["id"], now()),
                    )
                    cur.execute(
                        "UPDATE pipelines SET next_run = %s WHERE id = %s",
                        (
                            time.time() + (p["interval_minutes"] or 0) * 60,
                            p["id"],
                        ),
                    )
            except Exception:
                return
            finally:
                conn.close()
            self.store.audit(None, "scheduled_run", new_id)
            break

    def tick(self):
        """Atiende un run encolado. Devuelve True si hizo algo."""
        self.schedule()
        row = self.store.one(
            "SELECT id FROM runs WHERE status = 'queued'" " ORDER BY created_at LIMIT 1"
        )
        if not row:
            return False
        run_id = row["id"]
        conn = mydb.connect()
        try:
            with conn.cursor() as cur:
                cur.execute(
                    "UPDATE runs SET status = 'running', started_at = %s,"
                    " attempts = attempts + 1 WHERE id = %s AND status = 'queued'",
                    (now(), run_id),
                )
                if cur.rowcount == 0:
                    return False
        finally:
            conn.close()
        try:
            self.engine.process(run_id)
        except Exception as exc:  # noqa: BLE001 (worker nunca debe morir)
            import logging
            import pymysql as _pm

            info = self.store.one("SELECT attempts FROM runs WHERE id = %s", (run_id,))
            retry = isinstance(exc, (_pm.OperationalError, OSError)) and (
                info["attempts"] < 2
            )
            msg = (
                "Error temporal; se reintenta una vez."
                if retry
                else "Fallo el proceso. Revisa el log."
            )
            conn = mydb.connect()
            try:
                with conn.cursor() as cur:
                    cur.execute(
                        "UPDATE runs SET status = %s, stage = %s, error = %s,"
                        " finished_at = %s WHERE id = %s",
                        (
                            "queued" if retry else "failed",
                            "queued" if retry else "failed",
                            msg,
                            None if retry else now(),
                            run_id,
                        ),
                    )
            finally:
                conn.close()
            self.store.log(run_id, "queued" if retry else "failed", msg)
            if not retry:
                try:
                    rec = self.store.one(
                        "SELECT r.dataset_id, d.project_id FROM runs r"
                        " LEFT JOIN datasets d ON d.id = r.dataset_id"
                        " WHERE r.id = %s",
                        (run_id,),
                    )
                    self.store.alert(
                        "error",
                        "Pipeline fallido",
                        msg,
                        "run",
                        run_id,
                        (rec and rec.get("project_id")) or "p_default",
                    )
                except Exception:
                    pass
            logging.exception("Fallo run %s", run_id)
        return True

    def loop(self):
        while not self.stop_event.is_set():
            try:
                self.tick()
            except Exception:  # noqa: BLE001
                import logging

                logging.exception("Error en worker")
            self.stop_event.wait(10)

    def start(self):
        self.recover()
        self.thread = threading.Thread(
            target=self.loop, name="dataops-worker", daemon=True
        )
        self.thread.start()

    def stop(self):
        self.stop_event.set()
        if self.thread:
            self.thread.join(timeout=30)


# ---------------------------------------------------------------- export


def csv_export(rows):
    """Exporta filas a CSV con columnas dinámicas + estado/motivo/score."""
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    cols, decoded = [], []
    for row in rows:
        d = json.loads(row["data_json"])
        decoded.append((row, d))
        for k in d:
            if k not in cols:
                cols.append(k)
    writer.writerow((cols or list(COLUMNS)) + ["estado", "motivo", "score_ia"])
    for row, d in decoded:
        vals = []
        for v in [d.get(k, "") for k in cols] + [
            row["status"],
            row["reason"],
            row["score"],
        ]:
            v = "" if v is None else str(v)
            if v.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
                v = "'" + v
            vals.append(v)
        writer.writerow(vals)
    return "\ufeff" + out.getvalue()
