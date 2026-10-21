"""Lectura de archivos (CSV, Excel, SQL) y exportacion a CSV.

Normaliza cualquier archivo a un texto CSV con las columnas detectadas.
"""

import csv
import io
import json
import re

from backend.settings import COLUMNS, MAX_BYTES, MAX_ROWS, clean_header

# ---------------------------------------------------------------- parseo


def parse_csv_dynamic(content):
    """Lee CSV (coma o punto y coma) con cualquier columna.

    Devuelve (columns, rows). Falla si: vacío, >10 MiB, binario,
    cabecera inválida, filas disparejas, >20000 filas.
    """
    if not isinstance(content, str) or not content.strip():
        raise ValueError("El archivo está vacío.")
    if len(content.encode("utf-8")) > MAX_BYTES:
        raise ValueError("El tamaño máximo es 10 MiB.")
    if "\x00" in content:
        raise ValueError("El archivo tiene caracteres binarios. Usa UTF-8.")

    content = content.lstrip("\ufeff")
    header = content.splitlines()[0]
    delimiter = ";" if header.count(";") > header.count(",") else ","

    try:
        reader = csv.DictReader(
            io.StringIO(content, newline=""), delimiter=delimiter, strict=True
        )
        raw = reader.fieldnames or []
        columns = [clean_header(c) for c in raw]
        if (
            not columns
            or len(columns) > 30
            or any(not c for c in columns)
            or len(set(columns)) != len(columns)
        ):
            raise ValueError("Cabecera inválida: 1 a 30 columnas únicas.")

        rows = []
        for item in reader:
            if len(rows) >= MAX_ROWS:
                raise ValueError("El límite es de 20000 registros.")
            if None in item or any(v is None for v in item.values()):
                raise ValueError("Hay filas con columnas de más o de menos.")
            row = {
                col: str(item[orig]).strip() if item[orig] is not None else ""
                for col, orig in zip(columns, raw)
            }
            if any(len(v) > 2000 for v in row.values()):
                raise ValueError("Cada celda admite hasta 2000 caracteres.")
            rows.append(row)
    except csv.Error as exc:
        raise ValueError("CSV con comillas o separadores incorrectos.") from exc

    if not rows:
        raise ValueError("El archivo debe tener al menos un registro.")
    return columns, rows


def parse_csv_strict(content):
    """Valida que el CSV sea exactamente la plantilla de ventas."""
    columns, rows = parse_csv_dynamic(content)
    if columns != COLUMNS:
        raise ValueError(
            "Para ventas, las columnas deben ser: " + ", ".join(COLUMNS) + "."
        )
    return rows


def parse_excel_dynamic(data):
    """Lee XLSX (primera hoja) con cualquier columna."""
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise ValueError("El Excel está vacío.")
    if len(data) > MAX_BYTES:
        raise ValueError("El Excel supera los 10 MiB.")

    from openpyxl import load_workbook

    try:
        wb = load_workbook(
            filename=io.BytesIO(bytes(data)), read_only=True, data_only=True
        )
        values = list(wb.active.iter_rows(values_only=True))
    except Exception as exc:
        raise ValueError("No se pudo leer el Excel.") from exc

    if not values or len(values) < 2:
        raise ValueError("El Excel necesita cabecera + 1 registro mínimo.")

    raw = [str(c).strip() if c is not None else "" for c in values[0]]
    columns = [clean_header(c) for c in raw]
    if (
        not columns
        or len(columns) > 30
        or any(not c for c in columns)
        or len(set(columns)) != len(columns)
    ):
        raise ValueError("Cabecera Excel inválida: 1 a 30 columnas únicas.")

    rows = []
    for line in values[1:]:
        if all(c is None or str(c).strip() == "" for c in line):
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError("El límite es de 20000 registros.")
        cells = [("" if c is None else str(c).strip()) for c in list(line)]
        while len(cells) < len(columns):
            cells.append("")
        rows.append(dict(zip(columns, cells[: len(columns)])))

    if not rows:
        raise ValueError("El archivo debe tener al menos un registro.")
    return columns, rows


def parse_sql_dynamic(content):
    """Lee .sql con INSERT INTO tabla (col1, col2) VALUES (...).

    Acepta cualquier columna. Ignora comentarios -- y /* */.
    """
    if not isinstance(content, str) or not content.strip():
        raise ValueError("El SQL está vacío.")
    if len(content.encode("utf-8")) > MAX_BYTES:
        raise ValueError("El tamaño máximo es 10 MiB.")

    text = re.sub(r"/\*.*?\*/", " ", content, flags=re.DOTALL)
    kept = []
    for line in text.splitlines():
        s = line.strip()
        if not s or s.startswith("--"):
            continue
        kept.append(re.split(r"\s--\s", line)[0])
    text = " ".join(kept)

    inserts = list(
        re.finditer(
            r"INSERT\s+INTO\s+\S+\s*\(([^)]+)\)\s*VALUES\s*(.+?);",
            text,
            flags=re.IGNORECASE | re.DOTALL,
        )
    )
    if not inserts:
        raise ValueError(
            "Sin INSERTs. Usa: INSERT INTO tabla (col1, col2) VALUES (...);"
        )

    def split_cols(s):
        return [
            clean_header(c.strip().strip('"').strip("'").strip("`").strip("[]"))
            for c in s.split(",")
        ]

    columns = split_cols(inserts[0].group(1))
    if not columns or len(columns) > 30 or len(set(columns)) != len(columns):
        raise ValueError("Columnas SQL inválidas: 1 a 30 únicas.")

    token = re.compile(r"""'((?:[^']|'')*)'|"((?:[^"]|"")*)"|([^,()]+)""")
    rows = []
    for ins in inserts:
        cols_here = split_cols(ins.group(1))
        for m in re.finditer(r"\(([^;()]+)\)", ins.group(2)):
            cells = []
            for tok in token.finditer(m.group(1)):
                if tok.group(1) is not None:
                    cells.append(tok.group(1).replace("''", "'").strip())
                elif tok.group(2) is not None:
                    cells.append(tok.group(2).replace('""', '"').strip())
                else:
                    cells.append((tok.group(3) or "").strip())
            if len(cells) != len(cols_here):
                continue
            row = dict(zip(cols_here, cells))
            rows.append({c: row.get(c, "") for c in columns})
            if len(rows) >= MAX_ROWS:
                raise ValueError("El límite es de 20000 registros.")

    if not rows:
        raise ValueError("El SQL no aportó filas válidas.")
    return columns, rows


def rows_to_csv(columns, rows):
    """Convierte filas a texto CSV (formato interno de storage)."""
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=columns)
    writer.writeheader()
    for r in rows:
        writer.writerow({k: r.get(k, "") for k in columns})
    return out.getvalue()


# ---------------------------------------------------------------- export


def csv_export(rows):
    """Exporta filas a CSV con columnas dinámicas + estado/motivo/score."""
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    cols, decoded = [], []
    for row in rows:
        d = json.loads(row["data_json"])
        decoded.append((row, d))
        for k in d:
            if k not in cols:
                cols.append(k)
    writer.writerow((cols or list(COLUMNS)) + ["estado", "motivo", "score_ia"])
    for row, d in decoded:
        vals = []
        for v in [d.get(k, "") for k in cols] + [
            row["status"],
            row["reason"],
            row["score"],
        ]:
            v = "" if v is None else str(v)
            if v.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
                v = "'" + v
            vals.append(v)
        writer.writerow(vals)
    return "\ufeff" + out.getvalue()
